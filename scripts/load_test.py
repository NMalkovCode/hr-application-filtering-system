"""Developer load test: a full pipeline run over many generated applications with a simulated
Claude that is slow and sometimes fails. No API calls, no cost.

    .venv\\Scripts\\python.exe scripts\\load_test.py [--count 200] [--concurrency 4]

Checks that every application is accounted for, transient failures are retried, the output tree
and Summary.xlsx are complete, and reports memory use and run time.
"""

from __future__ import annotations

import argparse
import hashlib
import random
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import tracemalloc
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from hr_screening.careers import _RoleSummary  # noqa: E402
from hr_screening.evaluation import CandidateEvaluation, Judgement  # noqa: E402
from hr_screening.llm import ClaudeClient, RateLimited, Refused  # noqa: E402
from hr_screening.models import Tier  # noqa: E402
from hr_screening.pipeline import RunRequest, ScreeningPipeline  # noqa: E402
from hr_screening.settings import TierThresholds  # noqa: E402
from tests.fakes import fake_message, install_fake  # noqa: E402
from tests.helpers import evaluation, filter_set  # noqa: E402


def simulated_claude(failure_rate: float, refusal_rate: float):
    attempts: dict[str, int] = {}
    lock = threading.Lock()
    stats = {"rate_limited": 0, "refused": 0, "ok": 0}

    def respond(kwargs):
        time.sleep(random.uniform(0.05, 0.15))
        if kwargs["output_format"] is _RoleSummary:
            title = kwargs["messages"][0]["content"][0]["text"].split(">\n", 1)[1].splitlines()[0][:60]
            return fake_message(_RoleSummary(is_job_posting="Senior Data Analyst" not in title,
                                             title=title, summary="Summary."))
        text = " ".join(b.get("text", "") for b in kwargs["messages"][0]["content"])
        key = hashlib.md5(text.encode()).hexdigest()
        rng = random.Random(key)
        with lock:
            attempts[key] = attempts.get(key, 0) + 1
            first = attempts[key] == 1
        if rng.random() < refusal_rate:
            stats["refused"] += 1
            return Refused("declined")
        if first and rng.random() < failure_rate:
            stats["rate_limited"] += 1
            return RateLimited("rate limited")
        stats["ok"] += 1
        results = [rng.choices(list(Judgement), weights=[6, 2, 1, 1])[0] for _ in range(4)]
        alt = dict(alt_id="R01", alt_fit="strong", alt_better=rng.random() < 0.3) if rng.random() < 0.15 else {}
        return fake_message(evaluation(results, candidate_name=f"Person {key[:6]}", **alt),
                            input_tokens=3500, output_tokens=1800, cache_read=2000)

    return respond, stats


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--count", type=int, default=200)
    p.add_argument("--concurrency", type=int, default=4)
    args = p.parse_args()

    work = Path(tempfile.mkdtemp(prefix="hr_load_"))
    try:
        print(f"Generating {args.count} applications…")
        subprocess.run([sys.executable, str(ROOT / "scripts" / "make_trial_data.py"), "--count", str(args.count),
                        "--out", str(work / "data"), "--seed", "11"], check=True, stdout=subprocess.DEVNULL)

        client = ClaudeClient("sk-test")
        respond, stats = simulated_claude(failure_rate=0.10, refusal_rate=0.02)
        install_fake(client, respond)
        pipeline = ScreeningPipeline(client)
        request = RunRequest(applications_dir=work / "data" / "applications", output_dir=work / "out",
                             filter_set=filter_set(), thresholds=TierThresholds(80, 65, 50),
                             careers_folder=str(work / "data" / "open_roles"), concurrency=args.concurrency)

        tracemalloc.start()
        t0 = time.perf_counter()
        prep = pipeline.prepare(request)
        prep.evaluator.retry_delays = (0.2, 0.5)          # real runs wait 20s / 60s / 120s
        t1 = time.perf_counter()
        progress = []
        outcome = pipeline.run(prep, on_progress=lambda pr: progress.append(pr.done))
        t2 = time.perf_counter()
        _, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()

        counts = {t.value: outcome.count(t) for t in Tier}
        total = sum(counts.values())
        rows = load_workbook(outcome.summary_path)["Candidates"].max_row - 1
        reports = len(list(outcome.run_dir.rglob("*_Report.docx")))
        shortlisted = sum(outcome.count(t) for t in (Tier.TOP, Tier.CLOSE, Tier.BACKUP))
        needs_review = [p.result.error for p in outcome.placements if p.result.tier is Tier.NEEDS_REVIEW]

        print(f"\nPrepare: {t1 - t0:.1f}s   Evaluate + outputs: {t2 - t1:.1f}s   Peak Python memory: {peak / 1e6:.0f} MB")
        print(f"Open roles compared: {[r.title for r in prep.roles.roles]}")
        print(f"Simulated API: {stats}")
        print(f"Tiers: {counts}   Other Role copies: {outcome.other_role_count}")
        print(f"Needs Review reasons: { {r: needs_review.count(r) for r in set(needs_review)} }")
        checks = {
            "every application accounted for": total == len(prep.candidates) == rows,
            "progress reported for each": progress == list(range(1, len(prep.candidates) + 1)),
            "one report per shortlisted + Other Role": reports == shortlisted + outcome.other_role_count,
            "rate-limited calls recovered": not any("rate limited" in r for r in needs_review),
            "screened role excluded from open roles": all("Senior Data Analyst" not in r.title for r in prep.roles.roles),
            "no output errors": not any(pl.errors for pl in outcome.placements),
        }
        for name, ok in checks.items():
            print(f"  [{'PASS' if ok else 'FAIL'}] {name}")
        return 0 if all(checks.values()) else 1
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
