"""Generate docs/HR_User_Guide.docx (and .pdf when Microsoft Word is available).

    .venv\\Scripts\\python.exe packaging\\make_user_guide.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from hr_screening import __version__  # noqa: E402

SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "docs" / "HR_User_Guide.docx"
BLUE = RGBColor(0x1F, 0x5F, 0xA8)
GREY = RGBColor(0x59, 0x59, 0x59)


# --- helpers -----------------------------------------------------------------------

def setup(doc: Document) -> None:
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin, s.bottom_margin = Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, colour in (("Heading 1", 16, BLUE), ("Heading 2", 12.5, RGBColor(0x26, 0x26, 0x26))):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.color.rgb, st.font.bold = "Calibri", Pt(size), colour, True
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
    for list_style in ("List Bullet", "List Number"):
        doc.styles[list_style].font.size = Pt(10.5)
        doc.styles[list_style].paragraph_format.space_after = Pt(2)
    footer = s.footer.paragraphs[0]
    footer.text = f"HR Candidate Screening {__version__} — User Guide"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = GREY


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def para(doc, text, *, bold_prefix: str = "", italic=False, colour=None, size=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic = italic
    if colour:
        r.font.color.rgb = colour
    if size:
        r.font.size = Pt(size)
    return p


def bullets(doc, items, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(style=style)
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)


def table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        shade(cell, "E7ECF2")
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            cell.paragraphs[0].add_run(text.removeprefix(" — "))
    for i, row in enumerate(t.rows):
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        if i == 0:   # repeat the header row if a table crosses a page
            header_flag = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header_flag)
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                # Keep short tables on one page (together with their heading).
                p.paragraph_format.keep_with_next = i < len(t.rows) - 1
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def note(doc, text, fill="FFF8E1"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.6)
    shade(cell, fill)
    cell.text = ""
    cell.paragraphs[0].add_run(text).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def screenshot(doc, name, caption, width_cm=14.5):
    path = SHOTS / name
    if not path.exists():
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)
    c.runs[0].font.color.rgb = GREY


# --- content -----------------------------------------------------------------------

def build() -> Document:
    doc = Document()
    setup(doc)

    t = doc.add_paragraph()
    r = t.add_run("HR Candidate Screening")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), BLUE
    para(doc, f"User guide for HR teams · version {__version__}", colour=GREY, size=12)
    para(doc, "HR Candidate Screening reads a job description, proposes screening filters for you to check, "
              "then assesses a folder of applications against them. It sorts every candidate into a folder, "
              "ranks the best matches, writes a one-page report explaining each decision, and points out "
              "candidates who would suit another open role better. You stay in control: you verify the "
              "filters, and every placement is a recommendation for a person to review.")

    h1(doc, "1. Before you start")
    bullets(doc, [
        ("Windows 10 or 11 ", "and an internet connection."),
        ("Your own Anthropic API key. ", "Each HR user has their own key (ask IT or your team lead). The key "
         "is stored securely for your Windows account only."),
        ("Data protection. ", "Application content is sent to Anthropic's AI service for assessment. Make sure "
         "your organisation has approved this and that your candidate privacy notice covers it."),
    ])
    h2(doc, "Starting the program")
    bullets(doc, [
        "Double-click HRScreening.exe. No installation is needed; it can live on your desktop or a shared drive.",
        "The first start takes a few seconds while the program unpacks itself.",
        "If Windows shows \"Windows protected your PC\", ask IT whether the program is approved, then choose "
        "More info → Run anyway.",
        "On first launch the Settings window opens. Paste your API key, press Test key, then Save.",
    ], style="List Number")
    screenshot(doc, "6_settings.png", "Settings: API key, AI model, spending limit and Clear Cache", 9.5)

    h1(doc, "2. Screening applications step by step")
    para(doc, "The bar at the top of the window shows where you are: Job description → Filters → Applications "
              "→ Screening → Results.")

    h2(doc, "Step 1 — Job description")
    bullets(doc, [
        "Paste the job description, or press Load file… to open a PDF, Word document or scanned image.",
        "Press Analyse job description. This usually takes under a minute.",
        "If you have screened this job description before, the program offers the filters you verified "
        "last time — no new analysis needed. Use saved filters… lists every saved role.",
    ])
    screenshot(doc, "1_job_description.png", "Step 1: paste or load the job description")

    h2(doc, "Step 2 — Check the filters, then Verify")
    para(doc, "The program proposes a list of requirements (filters). This is the most important step: every "
              "application is judged against exactly these filters.")
    table(doc, ["Column", "What it means"], [
        ("Requirement", " — something a CV can show, e.g. \"At least 5 years' experience with SQL\". Edit the wording freely."),
        ("Must-have / Preferred", " — a candidate who does not meet a Must-have is not shortlisted. Keep Must-haves to "
                                  "true deal-breakers (usually 2–6)."),
        ("Weight 1–5", " — how much the requirement counts towards the score (5 = central to the job)."),
        ("Why", " — the part of the job description the filter came from, so you can check it."),
    ], [4.2, 12.4])
    bullets(doc, [
        "Use + Add filter, Remove selected and Move up / Move down to adjust the list (up to 25 filters).",
        "Read the yellow notes. They flag wording that was changed to avoid discrimination (for example "
        "\"native speaker\" becomes \"professional fluency\") and requirements better checked at interview.",
        "Press Verify filters. The filters are saved so the role can be screened again later.",
    ])
    screenshot(doc, "2_filters.png", "Step 2: review and edit the proposed filters")

    h2(doc, "Step 3 — Applications and options")
    bullets(doc, [
        ("Applications: ", "the folder that holds the applications. Best practice is one subfolder per candidate "
         "(CV, cover letter, certificates…). A folder with one file per candidate, or a mix, also works. PDF, "
         "Word (.docx) and scanned images (JPG, PNG, TIFF) are accepted. The originals are never changed."),
        ("Save results in: ", "where the results folder will be created. It must be outside the applications folder."),
        ("Other open roles (optional): ", "the address of your careers page, or a folder with one job description "
         "file per vacancy. Candidates who would suit another vacancy better are copied into an Other Role folder. "
         "Choose Don't check other roles to skip this."),
        ("Anonymised mode: ", "folders, reports and the spreadsheet use Candidate 001, 002… instead of names. A "
         "separate Candidate_Key.xlsx links the IDs back to the applications — keep it restricted."),
        ("Minimum scores: ", "the score needed for Top Match, Close Match and Backup (defaults 80 / 65 / 50)."),
    ])
    para(doc, "Press Start screening.")
    screenshot(doc, "3_setup.png", "Step 3: folders, other open roles, anonymised mode and score thresholds")

    h2(doc, "Step 4 — Screening")
    bullets(doc, [
        "The program reads the files, then assesses several applications at a time. Expect roughly 15–30 "
        "seconds per application, so 200 applications take around 15–35 minutes. Your PC will not go to "
        "sleep during a run.",
        "If the estimated cost could exceed your spending limit, you are asked before anything is spent.",
        "Cancel stops the run. Applications already assessed are saved; the rest go to Needs Review and the "
        "results folder name ends in _INCOMPLETE.",
    ])

    h2(doc, "Step 5 — Results")
    bullets(doc, [
        "The cards show how many candidates are in each folder.",
        "Double-click a candidate to open their report.",
        "Open results folder and Open Summary.xlsx take you to the outputs. Screen more applications for this "
        "role keeps the same verified filters.",
    ])
    screenshot(doc, "5_results.png", "Step 5: results overview", 14)

    h1(doc, "3. Understanding the results")
    table(doc, ["Folder", "Who is in it", "Report"], [
        ("Top Match", " — meets every must-have and scores at or above the Top Match threshold", "Yes"),
        ("Close Match", " — meets every must-have, scores between the Close Match and Top Match thresholds", "Yes"),
        ("Backup", " — meets every must-have, lower score", "Yes"),
        ("Other Role", " — clearly better suited to another open vacancy (shortlisted candidates appear here too, "
                       "as a copy)", "Yes"),
        ("Not Suitable", " — misses a must-have or scores below the Backup threshold", "No"),
        ("Needs Review", " — could not be assessed (unreadable file, not an application, AI error), or would not be "
                         "shortlisted but some of their files could not be read — check by hand", "No"),
    ], [3.2, 11.4, 2.0])
    para(doc, "Files are numbered in rank order (01_, 02_ …) inside each folder, so the best candidate is always "
              "at the top.")

    h2(doc, "How the score works")
    para(doc, "For each filter the AI records Met, Partially met, Not met or No evidence, quoting the application "
              "as evidence. The program then calculates the score itself: each filter earns its weight if met and "
              "half its weight if partially met, as a percentage of the total weight. Anyone who does not meet a "
              "Must-have is not shortlisted, whatever their score. When scores are equal, the AI's overall-fit "
              "rating (1–10, shown in Summary.xlsx) decides the order — for example, more years of directly "
              "relevant experience. A candidate is never rejected on a partial application: if some of their "
              "files could not be read, they go to Needs Review instead of Not Suitable.")

    h2(doc, "The one-page reports")
    bullets(doc, [
        "The first line notes any other role the candidate suits (Other Role reports instead state whether the "
        "candidate is also suitable for the role you screened).",
        "Header: tier, score, rank and how many must-haves are fully met.",
        "A short verdict, key strengths, gaps and risks, and a checklist of every filter with its evidence.",
        "Integrity notes appear when an application tried to influence the screening (for example hidden text "
        "saying \"rank me first\") or contains clear contradictions. Such text is ignored in the assessment.",
    ])

    h2(doc, "Summary.xlsx")
    table(doc, ["Sheet", "Contents"], [
        ("Candidates", " — everyone, with tier, rank, score, a one-line reason, links to reports and folders, "
                       "and whether they strongly fit another role."),
        ("Summary", " — counts per folder, how many shortlisted candidates also fit another role, and a breakdown "
                    "per vacancy; run details and the actual cost."),
        ("Filter Matrix", " — every candidate against every filter, colour-coded."),
        ("Filters", " — the verified filters used for this run."),
    ], [3.4, 13.2])

    h1(doc, "4. Costs")
    para(doc, "Each assessment is a paid call to the AI service, charged to the API key. With the default model "
              "(Claude Opus 5.5) a typical application costs a few cents; 200 applications typically cost $10–20. "
              "Sonnet 5.5 costs about half as much and Haiku 5.5 a small fraction, with somewhat less nuanced "
              "judgements. Set a Spending limit per run in Settings: you are warned before a run that could "
              "exceed it. The actual cost of every run is shown on the results screen and in Summary.xlsx.")

    h1(doc, "5. Fair and responsible use")
    bullets(doc, [
        "The tool supports decisions; it does not make them. Review shortlists, and check Not Suitable and "
        "Needs Review before rejecting anyone.",
        "Filters and assessments ignore protected characteristics (age, sex, nationality, ethnicity, religion, "
        "disability, marital or family status, and so on) and their proxies, and do not penalise career breaks.",
        "Use Anonymised mode when you want reviewers to assess without names.",
        "Results folders contain personal data. Store them securely and delete them in line with your "
        "retention policy.",
    ])

    h1(doc, "6. Troubleshooting")
    table(doc, ["Message or problem", "What to do"], [
        ("\"The API key was rejected\"", "Open Settings, paste the key again and press Test key."),
        ("\"Insufficient credit\"", "Ask whoever manages your Anthropic account to top it up."),
        ("\"Rate limit\" / many Needs Review results", "The program already waits and retries. If it keeps "
         "happening, lower Parallel requests in Settings to 2."),
        ("\"Could not reach the AI service\"", "Check the internet connection or ask IT about the proxy/firewall."),
        ("Other roles not available", "The careers page may build its job list with JavaScript. Save the job "
         "descriptions into a folder and choose Folder of job descriptions instead."),
        ("A file is \"password-protected\" or \"old .doc format\"", "Ask the candidate for a PDF or .docx, or "
         "re-save the file, then screen that application again."),
        ("Something else", "The log file is at %APPDATA%\\HRScreeningTool\\logs\\hr_screening.log — send it to "
         "IT. It never contains application content."),
    ], [6.0, 10.6])

    h1(doc, "7. Settings reference")
    table(doc, ["Setting", "Meaning"], [
        ("API key", "Your personal key, stored in Windows Credential Manager. Test key checks it at no cost."),
        ("AI model", "Opus 5.5 (best quality, default), Sonnet 5.5 (balanced) or Haiku 5.5 (lowest cost)."),
        ("Spending limit per run", "Warn before a run whose estimated cost could exceed this. 0 = no limit."),
        ("Parallel requests", "How many applications are assessed at once (default 4)."),
        ("Default output folder", "Suggested location for results."),
        ("Clear Cache", "Deletes saved filters and the cached list of other open roles. Results are not affected."),
    ], [4.6, 12.0])
    return doc


def main() -> int:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(str(OUT))
    print(f"Wrote {OUT}")
    try:
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        d = word.Documents.Open(str(OUT), ReadOnly=True)
        pages = d.ComputeStatistics(2)
        d.SaveAs2(str(OUT.with_suffix(".pdf")), FileFormat=17)
        d.Close(False)
        word.Quit()
        print(f"Wrote {OUT.with_suffix('.pdf')} ({pages} pages)")
    except Exception as exc:
        print(f"PDF not created (Microsoft Word needed): {exc}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
