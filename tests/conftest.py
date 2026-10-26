"""Fixtures that generate realistic sample application files on the fly."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF
from PIL import Image, ImageDraw

CV_TEXT = (
    "Jane Doe - Senior Data Analyst\n"
    "Experience: 6 years of SQL, Python and Power BI at Acme Corp.\n"
    "Education: BSc Mathematics, University of Leeds.\n"
    "Skills: stakeholder management, dashboarding, statistics, A/B testing."
)


def make_text_pdf(path: Path, text: str = CV_TEXT, pages: int = 1) -> Path:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for _ in range(pages):
        pdf.add_page()
        pdf.multi_cell(0, 6, text)
    pdf.output(str(path))
    return path


def make_scan_image(size=(1240, 1754), text: str = "SCANNED CV - John Smith") -> Image.Image:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate([text] + CV_TEXT.splitlines()):
        draw.text((80, 80 + i * 40), line, fill="black")
    return img


def make_scanned_pdf(path: Path) -> Path:
    """A PDF whose only content is a page image (no text layer)."""
    buf = io.BytesIO()
    make_scan_image().save(buf, format="PNG")
    buf.seek(0)
    pdf = FPDF()
    pdf.add_page()
    pdf.image(buf, x=0, y=0, w=210)
    pdf.output(str(path))
    return path


def make_docx(path: Path, text: str = CV_TEXT, header: str = "jane.doe@example.com") -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = header
    for line in text.splitlines():
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "English, French"
    doc.save(str(path))
    return path


def make_image(path: Path, size=(1240, 1754), fmt: str | None = None) -> Path:
    make_scan_image(size).save(path, format=fmt)
    return path


@pytest.fixture
def subfolder_layout(tmp_path: Path) -> Path:
    """One subfolder per candidate, as HR would usually provide."""
    root = tmp_path / "applications"
    (root / "Jane_Doe").mkdir(parents=True)
    make_text_pdf(root / "Jane_Doe" / "Jane_Doe_CV.pdf")
    make_docx(root / "Jane_Doe" / "Cover Letter.docx", text="Dear Hiring Manager, " * 10)
    (root / "Jane_Doe" / "~$Cover Letter.docx").write_bytes(b"lock file")

    (root / "John Smith").mkdir()
    make_scanned_pdf(root / "John Smith" / "scan.pdf")

    (root / "Candidate 10").mkdir()
    make_image(root / "Candidate 10" / "cv_photo.jpg")

    (root / "Candidate 2").mkdir()
    make_docx(root / "Candidate 2" / "resume.docx")
    (root / "Candidate 2" / "old_letter.doc").write_bytes(b"\xd0\xcf\x11\xe0")

    (root / "Empty Folder").mkdir()
    return root


@pytest.fixture
def file_layout(tmp_path: Path) -> Path:
    """One file per candidate directly in the folder."""
    root = tmp_path / "flat"
    root.mkdir()
    make_text_pdf(root / "alice-cv.pdf")
    make_docx(root / "bob_resume.docx")
    make_image(root / "carol.png")
    (root / "notes.xlsx").write_bytes(b"PK")
    (root / "Thumbs.db").write_bytes(b"")
    return root
