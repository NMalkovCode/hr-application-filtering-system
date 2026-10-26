from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from hr_screening import reports
from hr_screening.evaluation import Judgement
from hr_screening.filters import FilterCategory, FilterSet, FilterType, ScreeningFilter
from hr_screening.models import Tier
from hr_screening.reports import initial_role_line, write_other_role_report, write_tier_report
from hr_screening.scoring import assign_tiers, score_evaluation
from hr_screening.settings import TierThresholds

from .helpers import OPEN_ROLES, M, N, P, X, candidate, evaluation, filter_set
from .word import word_page_counter

WORDY = ("demonstrable hands-on experience designing building and maintaining complex enterprise "
         "data platforms across multiple business units with measurable commercial outcomes and "
         "stakeholder sign-off at board level")


def words(n: int) -> str:
    return " ".join((WORDY.split() * 5)[:n])


def big_filter_set(n: int) -> FilterSet:
    return FilterSet(job_title="Principal Data Platform Architect", filters=[
        ScreeningFilter(id=f"F{i:02d}", criterion=f"{i}. {words(22)}", weight=3, category=FilterCategory.SKILL,
                        type=FilterType.MUST_HAVE if i <= 3 else FilterType.PREFERRED)
        for i in range(1, n + 1)
    ])


def worst_case(n_filters: int, **kw):
    fs = big_filter_set(n_filters)
    results = ([M] * 3) + [P, M, X, N] * 10
    e = evaluation(results[:n_filters], verdict=words(80), strengths=[words(35)] * 6, gaps=[words(35)] * 4,
                   integrity_notes=[words(30)] * 3, **kw)
    for j in e.judgements:
        j.evidence = words(30)
    r = score_evaluation(candidate("Alexandra Montgomery-Fitzwilliam"), e, fs, OPEN_ROLES)
    return assign_tiers([r], TierThresholds(top=30, close=20, backup=10))[0], fs


@pytest.fixture(scope="module")
def word():
    counter = word_page_counter()
    if counter is None:
        pytest.skip("Microsoft Word is not available for page counting")
    yield counter
    counter.close()


def _text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts.extend(c.text for row in t.rows for c in row.cells)
    return "\n".join(parts)


def test_tier_report_content(tmp_path: Path):
    r = score_evaluation(candidate(), evaluation([M, P, M, X], alt_id="R01", alt_fit="strong", alt_better=True,
                                                 gaps=["No logistics background"]), filter_set(), OPEN_ROLES)
    (r,) = assign_tiers([r], TierThresholds(80, 65, 50))
    assert r.tier is Tier.CLOSE
    path = write_tier_report(r, filter_set(), "Jane Doe", tmp_path / "r.docx")
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text]
    assert lines[0].startswith("Also a strong fit for Data Engineer")
    assert lines[1] == "Jane Doe — Data Analyst"
    assert lines[2].startswith("Close Match") and "Score 71/100" in lines[2] and "Must-haves fully met 1/2" in lines[2]
    assert "Why this is a close match" in lines and "Gaps and risks" in lines
    text = _text(path)
    assert "Partially met" in text and "evidence 2" in text and "No evidence" in text


def test_tier_report_rejects_non_shortlisted(tmp_path: Path):
    r = score_evaluation(candidate(), evaluation([N, M, M, M]), filter_set())
    (r,) = assign_tiers([r], TierThresholds())
    with pytest.raises(ValueError):
        write_tier_report(r, filter_set(), "x", tmp_path / "x.docx")


@pytest.mark.parametrize("results, expected", [
    ([M, M, M, M], "Also suitable for Data Analyst: Top Match (100/100) — see the Top Match folder."),
    ([N, M, M, M], "Not suitable for Data Analyst: does not meet must-have — At least 5 years of SQL."),
    ([M, P, X, X], "Not suitable for Data Analyst: scored 54/100, below the shortlist threshold."),
])
def test_other_role_first_line(results, expected):
    r = score_evaluation(candidate(), evaluation(results, alt_id="R01", alt_fit="strong", alt_better=True),
                         filter_set(), OPEN_ROLES)
    (r,) = assign_tiers([r], TierThresholds(80, 65, 60))
    assert initial_role_line(r, filter_set()) == expected


def test_other_role_report_content(tmp_path: Path):
    r = score_evaluation(candidate(), evaluation([N, M, M, M], alt_id="R01", alt_fit="strong",
                                                 alternative_summary="Great pipeline engineer."), filter_set(), OPEN_ROLES)
    (r,) = assign_tiers([r], TierThresholds())
    path = write_other_role_report(r, filter_set(), "Jane Doe", tmp_path / "o.docx")
    lines = [p.text for p in Document(str(path)).paragraphs if p.text]
    assert lines[0].startswith("Not suitable for Data Analyst")
    assert lines[1] == "Jane Doe — better suited to Data Engineer"
    assert "Great pipeline engineer." in lines and "Why they suit Data Engineer" in lines
    # No gaps were given, so the section falls back to the failed must-have.
    assert any(l.startswith("Does not meet must-have: At least 5 years of SQL") for l in lines)
    assert "Builds pipelines; Python, Spark, cloud." in lines


def test_layout_compacts_only_when_needed():
    typical = score_evaluation(candidate(), evaluation([M, P, M, X], strengths=["a" * 120] * 4, gaps=["b" * 120] * 3),
                               filter_set())
    assert reports._choose_layout(typical, filter_set(), "") == reports.LAYOUTS[0]
    big, fs_big = worst_case(25)
    assert reports._choose_layout(big, fs_big, "x" * 200) != reports.LAYOUTS[0]


def test_realistic_twelve_filter_report_keeps_evidence(word, tmp_path: Path):
    fs = FilterSet(job_title="Senior Data Analyst", filters=[
        ScreeningFilter(id=f"F{i:02d}", criterion=f"At least 3 years' experience with {words(7)}", weight=3)
        for i in range(1, 13)
    ])
    e = evaluation([M, P, M, X] * 3, verdict=words(45), strengths=[words(20)] * 4, gaps=[words(20)] * 2)
    for j in e.judgements:
        j.evidence = words(15)
    (r,) = assign_tiers([score_evaluation(candidate(), e, fs)], TierThresholds(top=60, close=40, backup=10))
    assert reports._choose_layout(r, fs, "").evidence_words > 0
    assert word.pages(write_tier_report(r, fs, "Jane Doe", tmp_path / "r.docx")) == 1


def test_word_limits():
    assert reports._words("one two three four", 3) == "one two three…"
    assert reports._words("  one   two ", 5) == "one two"


@pytest.mark.parametrize("n_filters", [4, 8, 12, 16, 20, 25])
def test_tier_report_fits_one_page_in_word(word, tmp_path: Path, n_filters):
    r, fs = worst_case(n_filters, alt_id="R01", alt_fit="strong", alt_better=True)
    path = write_tier_report(r, fs, r.name, tmp_path / f"tier_{n_filters}.docx")
    assert word.pages(path) == 1


def test_other_role_report_fits_one_page_in_word(word, tmp_path: Path):
    r, fs = worst_case(25, alt_id="R01", alt_fit="strong", alt_better=True,
                       alternative_summary=words(80), alternative_reasons=[words(35)] * 5)
    path = write_other_role_report(r, fs, r.name, tmp_path / "other.docx")
    assert word.pages(path) == 1
