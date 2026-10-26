"""End-to-end run with real files and a fake Claude."""

from __future__ import annotations

import threading
from pathlib import Path

from openpyxl import load_workbook

from hr_screening.careers import _RoleSummary
from hr_screening.evaluation import CandidateEvaluation
from hr_screening.llm import ClaudeClient
from hr_screening.models import Tier
from hr_screening.pipeline import RunRequest, ScreeningPipeline
from hr_screening.settings import TierThresholds

from .conftest import make_docx, make_image, make_text_pdf
from .fakes import fake_message, install_fake
from .helpers import M, N, P, X, evaluation, filter_set

PROFILES = {   # CV text marker -> judgements and alternative-role answer
    "PROFILE-TOP": ([M, M, M, M], {}),
    "PROFILE-CLOSE": ([M, P, M, P], {}),
    "PROFILE-ENGINEER": ([N, M, M, X], dict(alt_fit="strong")),
}


def fake_claude(kwargs):
    if kwargs["output_format"] is _RoleSummary:
        title = kwargs["messages"][0]["content"][0]["text"].split(">\n", 1)[1].splitlines()[0]
        return fake_message(_RoleSummary(is_job_posting=True, title=title, summary=f"{title} summary"))
    assert kwargs["output_format"] is CandidateEvaluation
    text = " ".join(b.get("text", "") for b in kwargs["messages"][0]["content"])
    for marker, (results, kw) in PROFILES.items():
        if marker in text:
            if kw:
                # Find the Data Engineer role id from the system prompt.
                system = kwargs["system"][0]["text"]
                kw = dict(kw, alt_id=system.split('title="Data Engineer"')[0].rsplit('id="', 1)[1][:3])
            return fake_message(evaluation(results, **kw))
    return fake_message(evaluation([X, X, X, X], is_application=False))


def build_inputs(tmp_path: Path):
    apps = tmp_path / "applications"
    (apps / "Alice Top").mkdir(parents=True)
    make_text_pdf(apps / "Alice Top" / "cv.pdf", "Alice PROFILE-TOP " * 20)
    make_docx(apps / "Alice Top" / "letter.docx", text="Dear team " * 20)
    (apps / "Bob Close").mkdir()
    make_docx(apps / "Bob Close" / "resume.docx", text="Bob PROFILE-CLOSE " * 20)
    make_text_pdf(apps / "carol_engineer.pdf", "Carol PROFILE-ENGINEER " * 20)
    make_image(apps / "dave_scan.png", size=(600, 800))           # fake Claude: not an application
    (apps / "Empty").mkdir()
    (apps / "eve.docx").write_bytes(b"corrupt")

    roles = tmp_path / "open roles"
    roles.mkdir()
    (roles / "engineer.txt").write_text("Data Engineer\nPipelines", encoding="utf-8")
    (roles / "analyst.txt").write_text("Data Analyst\nThe role being screened", encoding="utf-8")
    return apps, roles


def test_full_run(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))   # isolate the roles cache
    apps, roles = build_inputs(tmp_path)
    client = ClaudeClient("sk-test")
    install_fake(client, fake_claude)
    pipeline = ScreeningPipeline(client)
    request = RunRequest(applications_dir=apps, output_dir=tmp_path / "out", filter_set=filter_set(),
                         thresholds=TierThresholds(80, 65, 50), careers_folder=str(roles), concurrency=3)

    statuses = []
    prep = pipeline.prepare(request, on_status=lambda m, d, t: statuses.append(m))
    assert len(prep.candidates) == 5                                     # Empty folder skipped
    assert [r.title for r in prep.roles.roles] == ["Data Engineer"]       # screened role excluded
    assert prep.estimate.candidates == 4 and prep.estimate.high_usd > 0   # corrupt eve.docx not counted
    assert any("Reading applications (5/5)" in s for s in statuses)

    outcome = pipeline.run(prep)
    run = outcome.run_dir
    assert run.parent == tmp_path / "out" and run.name.startswith("Data_Analyst_")
    assert outcome.count(Tier.TOP) == 1 and outcome.count(Tier.CLOSE) == 1
    assert outcome.count(Tier.NOT_SUITABLE) == 1 and outcome.count(Tier.NEEDS_REVIEW) == 2
    assert outcome.other_role_count == 1
    assert (run / "Top Match" / "01_Alice_Top_Report.docx").exists()
    assert (run / "Close Match" / "01_Bob_Close_Report.docx").exists()
    assert (run / "Other Role" / "Jane_Doe_Other_Role_Report.docx").exists()   # extracted name for a file
    assert sorted(p.name for p in (run / "Needs Review").iterdir()) == ["01_dave_scan", "02_eve"]
    assert outcome.summary_path.exists() and outcome.key_path is None

    ws = load_workbook(outcome.summary_path)["Summary"]
    labels = {r[0].value: r[1].value for r in ws.iter_rows() if r[0].value}
    assert labels["Applications processed"] == 5
    assert labels["Other roles compared"].startswith("1 from")


def test_unreadable_careers_source_skips_check(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    apps, _ = build_inputs(tmp_path)
    client = ClaudeClient("sk-test")
    install_fake(client, fake_claude)
    request = RunRequest(applications_dir=apps, output_dir=tmp_path / "out", filter_set=filter_set(),
                         careers_folder=str(tmp_path / "nope"))
    prep = ScreeningPipeline(client).prepare(request)
    assert "not found" in prep.roles_error and prep.roles.roles == []
    outcome = ScreeningPipeline(client).run(prep)
    assert not (outcome.run_dir / "Other Role").exists()


def test_cancelled_run_is_marked_incomplete(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    apps, _ = build_inputs(tmp_path)
    client = ClaudeClient("sk-test")
    install_fake(client, fake_claude)
    pipeline = ScreeningPipeline(client)
    prep = pipeline.prepare(RunRequest(applications_dir=apps, output_dir=tmp_path / "out", filter_set=filter_set()))
    cancel = threading.Event()
    outcome = pipeline.run(prep, on_progress=lambda p: cancel.set(), cancel=cancel)
    assert outcome.cancelled and outcome.run_dir.name.endswith("_INCOMPLETE")
    assert outcome.count(Tier.NEEDS_REVIEW) >= 4


def test_anonymised_run_writes_key(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    apps, _ = build_inputs(tmp_path)
    client = ClaudeClient("sk-test")
    install_fake(client, fake_claude)
    pipeline = ScreeningPipeline(client)
    prep = pipeline.prepare(RunRequest(applications_dir=apps, output_dir=tmp_path / "out",
                                       filter_set=filter_set(), anonymised=True))
    outcome = pipeline.run(prep)
    assert outcome.key_path.exists()
    names = " ".join(p.name for p in outcome.run_dir.rglob("*"))
    assert "Alice" not in names and "Candidate_001" in names
