from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from hr_screening.settings import TierThresholds
from hr_screening.organizer import place_all
from hr_screening.spreadsheet import RunInfo, write_candidate_key, write_summary

from .helpers import OPEN_ROLES, filter_set
from .test_organizer import _tiered


def _run(tmp_path: Path, anonymised=False):
    ordered = _tiered(tmp_path)
    run = tmp_path / "run"
    run.mkdir()
    placements = place_all(ordered, filter_set(), run, anonymised=anonymised, other_roles_enabled=True)
    info = RunInfo(job_title="Data Analyst", started=datetime(2026, 10, 26, 9), finished=datetime(2026, 10, 26, 9, 20),
                   model="Claude Opus 5.5", thresholds=TierThresholds(80, 65, 50), anonymised=anonymised,
                   applications_dir=str(tmp_path / "apps"), cost_usd=1.234, other_roles_source="https://acme.test/careers",
                   open_roles=OPEN_ROLES, notes=["notes.xlsx: Unsupported file type"])
    path = write_summary(run / "Summary.xlsx", placements, filter_set(), info)
    return load_workbook(path), placements, run


def rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_candidates_sheet(tmp_path: Path):
    wb, _, run = _run(tmp_path)
    assert wb.sheetnames == ["Candidates", "Summary", "Filter Matrix", "Filters"]
    data = rows(wb["Candidates"])
    header, body = data[0], data[1:]
    col = {h: i for i, h in enumerate(header)}
    assert [r[col["Candidate"]] for r in body] == ["Amy Top", "Ben Top", "Cat Backup", "Dan Out", "Eve Broken"]
    ben = body[1]
    assert ben[col["Tier"]] == "Top Match" and ben[col["Strong fit for another role"]] == "Yes"
    assert ben[col["Other role"]] == "Data Engineer" and ben[col["Copied to Other Role"]] == "Yes"
    assert body[0][col["Must-haves fully met"]] == "2/2"
    assert body[0][col["Overall fit (1-10)"]] == 5 and body[4][col["Overall fit (1-10)"]] is None
    assert body[4][col["Summary"]] == "Unreadable" and body[4][col["Score"]] is None
    link = wb["Candidates"].cell(2, col["Report"] + 1).hyperlink.target
    assert link == "Top Match\\01_Amy_Top_Report.docx" and (run / link).exists()
    assert "Original location" in header


def test_summary_cross_tab_tracks_other_role_fits(tmp_path: Path):
    wb, _, _ = _run(tmp_path)
    data = rows(wb["Summary"])
    by_label = {r[0]: r[1:4] for r in data if r[0]}
    assert by_label["Top Match"] == [2, 1, 1]
    assert by_label["Not Suitable"] == [1, 1, 1]
    assert by_label["Shortlisted (Top + Close + Backup)"] == [3, 1, 1]
    assert by_label["Total"] == [5, 2, 2]
    assert by_label["Data Engineer"][:3] == [1, 1, 0]       # strong fits, of which Top Match, Close Match
    assert by_label["Actual AI cost (USD)"][0] == 1.23
    assert any(r[0] == "notes.xlsx: Unsupported file type" for r in data)


def test_filter_matrix(tmp_path: Path):
    wb, _, _ = _run(tmp_path)
    data = rows(wb["Filter Matrix"])
    assert data[0][4].startswith("F01 (Must, w5): At least 5 years of SQL")
    assert data[2][4:8] == ["Met", "Met", "Met", "Partially met"]


def test_anonymised_summary_and_key(tmp_path: Path):
    wb, placements, run = _run(tmp_path, anonymised=True)
    data = rows(wb["Candidates"])
    assert "Original location" not in data[0]
    flat = " ".join(str(v) for r in data for v in r if v)
    assert "Amy" not in flat and "Candidate 001" in flat

    key = rows(load_workbook(write_candidate_key(run / "Candidate_Key.xlsx", placements)).active)
    assert key[1][:2] == ["Candidate 001", "Jane Doe"] and key[1][2].endswith("Amy Top")
