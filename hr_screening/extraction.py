"""Read application files (PDF, DOCX, images, text) into SourceDocuments.

Text is extracted locally wherever possible. Scanned material (image-only PDFs, photos, scans)
is kept as bytes and later sent to Claude, which reads it directly - no OCR engine needed.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path

from docx import Document as load_docx
from docx.oxml.ns import qn
from PIL import Image, ImageOps, UnidentifiedImageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from . import config
from .models import ContentKind, SourceDocument

log = logging.getLogger(__name__)

_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
# A DOCX with less text than this but with embedded pictures is treated as a scanned document.
_MIN_DOCX_TEXT_CHARS = 80
_MIN_EMBEDDED_IMAGE_EDGE_PX = 400


class ExtractionError(Exception):
    """A file could not be read. The message is shown to HR as-is."""


def extract_document(path: Path) -> list[SourceDocument]:
    """Read one file. Returns one or more documents (multi-page TIFFs give one per page)."""
    ext = path.suffix.lower()
    try:
        size = path.stat().st_size
    except FileNotFoundError as exc:
        raise ExtractionError("File no longer exists.") from exc
    except OSError as exc:
        raise ExtractionError("File cannot be accessed.") from exc
    if size == 0:
        raise ExtractionError("File is empty.")
    try:
        if ext in config.PDF_EXTENSIONS:
            return [_extract_pdf(path)]
        if ext in config.DOCX_EXTENSIONS:
            return _extract_docx(path)
        if ext in config.IMAGE_EXTENSIONS:
            return _extract_image(path)
        if ext in config.TEXT_EXTENSIONS:
            return [_extract_text_file(path)]
    except ExtractionError:
        raise
    except PermissionError as exc:
        raise ExtractionError("File is open in another program or access is denied.") from exc
    except Exception as exc:  # corrupt files raise all sorts of library errors
        log.exception("Failed to read %s", path)
        raise ExtractionError(f"Could not read file ({type(exc).__name__}).") from exc
    raise ExtractionError(f"Unsupported file type '{ext}'.")


# --- PDF ---------------------------------------------------------------------------

def _extract_pdf(path: Path) -> SourceDocument:
    try:
        reader = PdfReader(path)
    except PdfReadError as exc:
        raise ExtractionError("PDF is damaged or not a valid PDF.") from exc

    if reader.is_encrypted:
        # Many PDFs are "encrypted" with an empty user password purely to restrict editing.
        try:
            if not reader.decrypt(""):
                raise ExtractionError("PDF is password-protected.")
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("PDF is password-protected.") from exc

    pages = [(page.extract_text() or "") for page in reader.pages]
    page_count = len(pages)
    if page_count == 0:
        raise ExtractionError("PDF has no pages.")

    text = _clean_text("\n\n".join(pages))
    if len(text) >= config.MIN_TEXT_CHARS_PER_PAGE * page_count:
        return SourceDocument(path=path, kind=ContentKind.TEXT, text=text, page_count=page_count)

    # Little or no text layer: a scanned PDF. Send the PDF itself to Claude.
    data = path.read_bytes()
    if len(data) > config.MAX_PDF_BYTES:
        raise ExtractionError(
            f"Scanned PDF is too large to process ({len(data) / 1e6:.0f} MB; "
            f"limit {config.MAX_PDF_BYTES / 1e6:.0f} MB)."
        )
    return SourceDocument(
        path=path,
        kind=ContentKind.PDF,
        text=text,  # whatever partial text exists, kept for cost estimates/debugging
        data=data,
        media_type="application/pdf",
        page_count=page_count,
    )


# --- DOCX --------------------------------------------------------------------------

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _extract_docx(path: Path) -> list[SourceDocument]:
    with open(path, "rb") as f:
        head = f.read(8)
    if head == _OLE_MAGIC:
        # Password-protected .docx files (and renamed .doc files) are OLE containers, not ZIPs.
        raise ExtractionError("Word document is password-protected or in the old .doc format.")
    try:
        doc = load_docx(str(path))
    except Exception as exc:
        raise ExtractionError("Word document is damaged or not a .docx file.") from exc

    parts: list[str] = []
    # Headers first: CV templates often put name and contact details there.
    # (Check is_linked_to_previous before touching _element: accessing it on a linked header
    # would create an empty definition.)
    seen_headers: list = []
    for section in doc.sections:
        for header in (section.header, section.first_page_header):
            if header.is_linked_to_previous:
                continue
            element = header._element
            if any(element is seen for seen in seen_headers):
                continue
            seen_headers.append(element)
            parts.extend(_paragraph_texts(element))

    # Body, in document order, including tables and text boxes.
    parts.extend(_paragraph_texts(doc.element.body))

    text = _clean_text("\n".join(parts))
    if len(text) >= _MIN_DOCX_TEXT_CHARS:
        return [SourceDocument(path=path, kind=ContentKind.TEXT, text=text)]

    # Almost no text: possibly a scan pasted into Word. Use the embedded pictures.
    images = _docx_embedded_images(doc, path)
    if images:
        if text:
            images.insert(0, SourceDocument(path=path, kind=ContentKind.TEXT, text=text))
        return images
    if text:
        return [SourceDocument(path=path, kind=ContentKind.TEXT, text=text)]
    raise ExtractionError("Word document contains no readable text.")


def _paragraph_texts(root) -> list[str]:
    """Text of every w:p under `root`, skipping duplicate VML fallback copies of text boxes."""
    texts = []
    for p in root.iter(qn("w:p")):
        if any(a.tag == _MC_FALLBACK for a in p.iterancestors()):
            continue
        chunks = []
        for node in p.iter():
            if node.tag == qn("w:t") and node.text:
                chunks.append(node.text)
            elif node.tag == qn("w:tab"):
                chunks.append("\t")
            elif node.tag in (qn("w:br"), qn("w:cr")):
                chunks.append("\n")
        line = "".join(chunks).strip()
        if line:
            texts.append(line)
    return texts


def _docx_embedded_images(doc, path: Path) -> list[SourceDocument]:
    docs = []
    for rel in doc.part.rels.values():
        if "image" not in rel.reltype or rel.is_external:
            continue
        blob = rel.target_part.blob
        try:
            if min(Image.open(io.BytesIO(blob)).size) < _MIN_EMBEDDED_IMAGE_EDGE_PX:
                continue  # logos, icons
        except (UnidentifiedImageError, OSError):
            continue  # e.g. EMF/WMF vector graphics
        try:
            docs.extend(_images_from_bytes(blob, path))
        except ExtractionError:
            continue
    return docs


# --- Images ------------------------------------------------------------------------

def _extract_image(path: Path) -> list[SourceDocument]:
    return _images_from_bytes(path.read_bytes(), path)


def _images_from_bytes(raw: bytes, path: Path) -> list[SourceDocument]:
    try:
        img = Image.open(io.BytesIO(raw))
    except (UnidentifiedImageError, OSError) as exc:
        raise ExtractionError("Image is damaged or in an unsupported format.") from exc

    frames = getattr(img, "n_frames", 1)
    is_gif = (img.format or "").upper() == "GIF"
    docs = []
    for index in range(1 if is_gif else frames):  # multi-page TIFF scans
        img.seek(index)
        docs.append(_prepare_image(img.copy(), raw if frames == 1 else None, path))
    return docs


def _prepare_image(img: Image.Image, original: bytes | None, path: Path) -> SourceDocument:
    """Return a Claude-ready image: correct orientation, long edge <= limit, under size cap."""
    native_type = config.NATIVE_IMAGE_TYPES.get(path.suffix.lower())
    needs_rotation = img.getexif().get(0x0112, 1) not in (0, 1)  # EXIF orientation tag
    fits = max(img.size) <= config.MAX_IMAGE_EDGE_PX

    if original is not None and native_type and fits and not needs_rotation \
            and len(original) <= config.MAX_IMAGE_BYTES:
        return SourceDocument(path=path, kind=ContentKind.IMAGE, data=original, media_type=native_type)

    img = ImageOps.exif_transpose(img)
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    if not fits:
        img.thumbnail((config.MAX_IMAGE_EDGE_PX, config.MAX_IMAGE_EDGE_PX), Image.Resampling.LANCZOS)

    # JPEG keeps photographed/scanned pages small; quality is high enough for small print.
    for quality in (90, 80, 70):
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        if buf.tell() <= config.MAX_IMAGE_BYTES:
            return SourceDocument(
                path=path, kind=ContentKind.IMAGE, data=buf.getvalue(), media_type="image/jpeg"
            )
    raise ExtractionError("Image is too large to process even after compression.")


# --- Plain text --------------------------------------------------------------------

def _extract_text_file(path: Path) -> SourceDocument:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    text = _clean_text(text)
    if not text:
        raise ExtractionError("Text file is empty.")
    return SourceDocument(path=path, kind=ContentKind.TEXT, text=text)


# --- helpers -----------------------------------------------------------------------

_MULTI_BLANK = re.compile(r"\n{3,}")
_TRAILING_WS = re.compile(r"[ \t]+\n")


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _TRAILING_WS.sub("\n", text)
    text = _MULTI_BLANK.sub("\n\n", text)
    return text.strip()
