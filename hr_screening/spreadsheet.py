"""Summary.xlsx (and Candidate_Key.xlsx in anonymised mode)."""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .evaluation import Judgement
from .filters import FilterSet, FilterType
from .models import OpenRole, Tier
from .organizer import Placement
from .settings import TierThresholds

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
TIER_FILLS = {
    Tier.TOP: "D9EAD3", Tier.CLOSE: "DCE6F2", Tier.BACKUP: "FCE8CC",
    Tier.NOT_SUITABLE: "F2F2F2", Tier.NEEDS_REVIEW: "F4CCCC",
}
RESULT_FILLS = {
    Judgement.MET: "D9EAD3", Judgement.PARTIAL: "FCE8CC",
    Judgement.NOT_MET: "F4CCCC", Judgement.NO_EVIDENCE: "F2F2F2",
}
TIER_ORDER = [Tier.TOP, Tier.CLOSE, Tier.BACKUP, Tier.NOT_SUITABLE, Tier.NEEDS_REVIEW]


@dataclass
class RunInfo:
    job_title: str
    started: datetime
    finished: datetime
    model: str
    thresholds: TierThresholds
    anonymised: bool
    applications_dir: str
    cost_usd: float
    other_roles_source: str = ""          # URL / folder, or "" when the check was skipped
    open_roles: list[OpenRole] = field(default_factory=list)
    cancelled: bool = False
    notes: list[str] = field(default_factory=list)


def write_summary(path: Path, placements: list[Placement], filter_set: FilterSet, info: RunInfo) -> Path:
    wb = Workbook()
    _candidates_sheet(wb.active, placements, info, Path(path).parent)
    _summary_sheet(wb.create_sheet("Summary"), placements, info)
    _matrix_sheet(wb.create_sheet("Filter Matrix"), placements, filter_set)
    _filters_sheet(wb.create_sheet("Filters"), filter_set)
    wb.save(path)
    return Path(path)


def write_candidate_key(path: Path, placements: list[Placement]) -> Path:
    """Maps pseudonyms back to the original applications - keep this file restricted."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidate Key"
    ws.append(["Pseudonym", "Name in application", "Original folder / file", "Tier"])
    for p in placements:
        c = p.result.candidate
        extracted = p.result.evaluation.candidate_name if p.result.evaluation else ""
        ws.append([p.pseudonym, extracted or c.display_name, str(c.source_path), p.result.tier.value])
    _style_table(ws, widths=[16, 28, 70, 16])
    wb.save(path)
    return Path(path)


# --- sheets ------------------------------------------------------------------------

def _candidates_sheet(ws: Worksheet, placements: list[Placement], info: RunInfo, base: Path):
    ws.title = "Candidates"
    roles_on = bool(info.other_roles_source)
    headers = ["Tier", "Rank", "Candidate", "Score", "Overall fit (1-10)", "Must-haves fully met",
               "Strong fit for another role", "Other role", "Copied to Other Role", "Summary", "Integrity notes",
               "Report", "Application"]
    if not info.anonymised:
        headers.append("Original location")
    ws.append(headers)
    col = {h: i for i, h in enumerate(headers, 1)}

    for p in placements:
        r = p.result
        scored_must = [s for s in r.scored if s.filter.type is FilterType.MUST_HAVE]
        must =(f"{sum(s.result is Judgement.MET for s in scored_must)}/{len(scored_must)}"
                if r.evaluation and scored_must else "")
        alt = r.alternative_role
        row = [
            r.tier.value, r.rank, p.name,
            r.score if r.evaluation else None,
            r.overall_fit if r.evaluation else None,
            must,
            ("Yes" if alt else "No") if roles_on and r.evaluation else "—",
            alt.title if alt else "",
            "Yes" if p.other_role_dir else "No" if roles_on else "—",
            r.one_line_reason,
            " ".join(r.evaluation.integrity_notes) if r.evaluation else "",
            "Open report" if p.report else "",
            "Open folder" if p.application_dir else "",
        ]
        if not info.anonymised:
            row.append(str(r.candidate.source_path))
        ws.append(row)
        row_idx = ws.max_row
        ws.cell(row_idx, 1).fill = PatternFill("solid", fgColor=TIER_FILLS[r.tier])
        _link(ws.cell(row_idx, col["Report"]), p.report, base)
        _link(ws.cell(row_idx, col["Application"]), p.application_dir, base)
        if p.errors:
            ws.cell(row_idx, col["Summary"]).value = f"{r.one_line_reason} [Output problem: {'; '.join(p.errors)}]"

    widths = [14, 6, 26, 7, 9, 11, 11, 24, 11, 60, 40, 13, 13] + ([60] if not info.anonymised else [])
    _style_table(ws, widths=widths, wrap_cols={col["Summary"], col["Integrity notes"]})


def _summary_sheet(ws: Worksheet, placements: list[Placement], info: RunInfo):
    ws["A1"] = f"Screening summary — {info.job_title}"
    ws["A1"].font = TITLE_FONT
    t = info.thresholds
    details = [
        ("Run started", info.started.strftime("%d %b %Y %H:%M")),
        ("Run finished", info.finished.strftime("%d %b %Y %H:%M") + ("  (CANCELLED — incomplete)" if info.cancelled else "")),
        ("Applications folder", info.applications_dir),
        ("Applications processed", len(placements)),
        ("AI model", info.model),
        ("Actual AI cost (USD)", round(info.cost_usd, 2)),
        ("Tier thresholds", f"Top Match ≥ {t.top}, Close Match ≥ {t.close}, Backup ≥ {t.backup}"),
        ("Anonymised mode", "On" if info.anonymised else "Off"),
        ("Other roles compared", f"{len(info.open_roles)} from {info.other_roles_source}" if info.other_roles_source
         else "Skipped (no careers page or folder given)"),
    ]
    row = 3
    for label, value in details:
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, value)
        row += 1

    # Tier counts, with how many in each tier also strongly fit another role.
    row += 1
    ws.cell(row, 1, "Results by tier").font = SECTION_FONT
    row += 1
    header_row = row
    for col, h in enumerate(["Tier", "Candidates", "Strong fit for another role", "Copied to Other Role"], 1):
        ws.cell(row, col, h)
    for tier in TIER_ORDER:
        group = [p for p in placements if p.result.tier is tier]
        row += 1
        ws.cell(row, 1, tier.value).fill = PatternFill("solid", fgColor=TIER_FILLS[tier])
        ws.cell(row, 2, len(group))
        ws.cell(row, 3, sum(1 for p in group if p.result.alternative_role))
        ws.cell(row, 4, sum(1 for p in group if p.other_role_dir))
    shortlisted = [p for p in placements if p.result.tier.is_shortlisted]
    for label, group in (("Shortlisted (Top + Close + Backup)", shortlisted), ("Total", placements)):
        row += 1
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, len(group)).font = Font(bold=True)
        ws.cell(row, 3, sum(1 for p in group if p.result.alternative_role)).font = Font(bold=True)
        ws.cell(row, 4, sum(1 for p in group if p.other_role_dir)).font = Font(bold=True)
    _header(ws, header_row, 4)
    _grid(ws, header_row, row, 4)

    # Per other role: how many candidates fit it, broken down by their tier for this role.
    if info.open_roles:
        row += 2
        ws.cell(row, 1, "Strong fits by other role").font = SECTION_FONT
        row += 1
        header_row = row
        cols = ["Other role", "Strong fits", *[t.value for t in TIER_ORDER[:4]], "Copied to Other Role"]
        for col, h in enumerate(cols, 1):
            ws.cell(row, col, h)
        for role in info.open_roles:
            fits = [p for p in placements if p.result.alternative_role and p.result.alternative_role.role_id == role.role_id]
            if not fits:
                continue
            row += 1
            ws.cell(row, 1, role.title)
            ws.cell(row, 2, len(fits))
            for col, tier in enumerate(TIER_ORDER[:4], 3):
                ws.cell(row, col, sum(1 for p in fits if p.result.tier is tier))
            ws.cell(row, 7, sum(1 for p in fits if p.other_role_dir))
        if row == header_row:
            row += 1
            ws.cell(row, 1, "No candidate was a strong fit for another open role.")
        _header(ws, header_row, len(cols))
        _grid(ws, header_row, row, len(cols))

    if info.notes:
        row += 2
        ws.cell(row, 1, "Notes").font = SECTION_FONT
        for note in info.notes:
            row += 1
            ws.cell(row, 1, note).alignment = Alignment(wrap_text=False)

    for col, width in enumerate([34, 18, 26, 22, 14, 14, 20], 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _matrix_sheet(ws: Worksheet, placements: list[Placement], filter_set: FilterSet):
    ws.append(["Tier", "Rank", "Candidate", "Score",
               *[f"{f.id} ({'Must' if f.type is FilterType.MUST_HAVE else 'Pref'}, w{f.weight}): {f.criterion}"
                 for f in filter_set.filters]])
    for p in placements:
        by_id = {s.filter.id: s.result for s in p.result.scored}
        ws.append([p.result.tier.value, p.result.rank, p.name, p.result.score if p.result.evaluation else None,
                   *[by_id[f.id].label if f.id in by_id else "" for f in filter_set.filters]])
        for col, f in enumerate(filter_set.filters, 5):
            if f.id in by_id:
                ws.cell(ws.max_row, col).fill = PatternFill("solid", fgColor=RESULT_FILLS[by_id[f.id]])
    _style_table(ws, widths=[14, 6, 26, 7] + [18] * len(filter_set.filters), header_height=75)


def _filters_sheet(ws: Worksheet, filter_set: FilterSet):
    ws.append(["ID", "Requirement", "Category", "Type", "Weight", "Reason (from the job description)"])
    for f in filter_set.filters:
        ws.append([f.id, f.criterion, f.category.label, f.type.label, f.weight, f.rationale])
    _style_table(ws, widths=[6, 60, 14, 11, 8, 60], wrap_cols={2, 6})


# --- formatting --------------------------------------------------------------------

def _style_table(ws: Worksheet, widths: list[int], wrap_cols: set[int] = frozenset(), header_height: float | None = None):
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    _header(ws, 1, ws.max_column)
    if header_height:
        ws.row_dimensions[1].height = header_height
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = WRAP if cell.column in wrap_cols else TOP
    ws.freeze_panes = "D2" if ws.max_column > 4 else "A2"
    ws.auto_filter.ref = ws.dimensions


def _header(ws: Worksheet, row: int, n_cols: int):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row, col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def _grid(ws: Worksheet, first: int, last: int, n_cols: int):
    for r in range(first, last + 1):
        for c in range(1, n_cols + 1):
            ws.cell(r, c).border = BORDER


def _link(cell, target: Path | None, base: Path):
    if target is None:
        return
    cell.hyperlink = os.path.relpath(target, base)
    cell.font = Font(color="0563C1", underline="single")
